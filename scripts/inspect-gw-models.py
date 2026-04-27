"""Inspect GW sheet to find duplicate model rows."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

FILE = Path(
    r"C:\Users\CHEIKH MBACKE DIENG\Documents\AUDIT&CONTROLE DE GESTION"
    r"\BUDGET 2026\REVISION BUDGET 2026 VENTILLE RIST NOU CMD.xlsx"
)

wb = openpyxl.load_workbook(FILE, read_only=False, data_only=True)
ws = wb["GW"]

print(f"GW sheet: {ws.max_row} rows × {ws.max_column} cols")
print()
print("All UNITES rows (model + row number):")

# Find all rows where col B = "UNITES"
for r in range(1, ws.max_row + 1):
    indic = ws.cell(row=r, column=2).value
    model_a = ws.cell(row=r, column=1).value
    if isinstance(indic, str) and indic.strip().upper() == "UNITES":
        print(f"  Row {r:3d}: A={model_a!r:<40} B=UNITES")

print()
print("Looking specifically for 'H6 TOP' rows:")
for r in range(1, ws.max_row + 1):
    val = ws.cell(row=r, column=1).value
    if isinstance(val, str) and "H6 TOP" in val.upper():
        print(f"  Row {r:3d}: A={val!r}  B={ws.cell(row=r, column=2).value!r}")
